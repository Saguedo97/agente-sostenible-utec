"""Consulta local indexada de CUIs 2023-2026. Sin etiquetas historicas.

Preparacion UNA vez: python universo_v2.py preparar "Data 23_26.xlsx"
Por seguridad, la base SQLite resultante NO se publica en GitHub.
"""
from __future__ import annotations

import argparse
import json
import os
import sqlite3
from contextlib import closing
from datetime import date, datetime
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = Path(os.getenv('UNIVERSE_DB_PATH', BASE_DIR / 'data' / 'universo_2023_2026.sqlite'))

# Campos de lectura: excluyen todo ground truth, emisiones y clasificaciones.
COLUMNAS = (
    'CODIGO_UNICO', 'NOMBRE_INVERSION', 'NIVEL', 'SECTOR', 'ENTIDAD',
    'ESTADO', 'SITUACION', 'FUNCION', 'PROGRAMA', 'SUBPROGRAMA',
    'TIPO_INVERSION', 'DES_TIPOLOGIA', 'DEPARTAMENTO', 'PROVINCIA',
    'DISTRITO', 'ALTERNATIVA', 'MONTO_VIABLE', 'COSTO_ACTUALIZADO',
    'EXPEDIENTE_TECNICO', 'ETAPA_F8', 'ETAPA_F9', 'PRIMER_DEVENGADO',
    'FEC_INI_EJECUCION', 'DEVENGADO_2023', 'DEVENGADO_2024',
    'DEVENGADO_2025', 'DEVENGADO_2026', 'TOTAL_DEV', 'TIENE_GASTO',
    'EMPEZO', 'NO_HA_EMPEZADO', 'N_BRECHAS', 'BRECHAS_TODAS',
)


def _valor(v):
    if isinstance(v, (date, datetime)):
        return v.isoformat()
    if isinstance(v, float) and (v != v or abs(v) == float('inf')):
        return None
    return v


def preparar_excel(origen: str | Path, destino: str | Path = DB_PATH) -> dict:
    """Lee XLSX en modo streaming para evitar cargar 121 mil filas en RAM."""
    from openpyxl import load_workbook

    destino = Path(destino)
    origen = Path(origen)
    if not origen.is_file():
        raise FileNotFoundError(origen)
    if origen.resolve() == destino.resolve():
        raise ValueError('Origen y destino deben ser diferentes.')
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(origen, read_only=True, data_only=True)
    try:
        hoja = wb['proyectos_terna_brechas'] if 'proyectos_terna_brechas' in wb else wb.active
        filas = hoja.iter_rows(values_only=True)
        encabezados = [str(v).strip() if v is not None else '' for v in next(filas)]
        if 'CODIGO_UNICO' not in encabezados:
            raise ValueError('El Excel no tiene la columna CODIGO_UNICO.')
        indices = {name: encabezados.index(name) for name in COLUMNAS if name in encabezados}
        temporal = destino.with_name(destino.name + '.tmp')
        if temporal.exists():
            temporal.unlink()
        conn = sqlite3.connect(temporal)
        total = 0
        try:
            conn.execute('CREATE TABLE proyectos (cui TEXT PRIMARY KEY, payload TEXT NOT NULL)')
            conn.execute('CREATE TABLE metadata (clave TEXT PRIMARY KEY, valor TEXT)')
            lote = []
            for fila in filas:
                d = {campo: _valor(fila[pos]) for campo, pos in indices.items()}
                v = d.get('CODIGO_UNICO')
                if v in (None, ''):
                    continue
                cui = str(v).strip().removesuffix('.0')
                if not cui.isdigit():
                    continue
                lote.append((cui, json.dumps(d, ensure_ascii=False, default=str)))
                if len(lote) >= 500:
                    conn.executemany('INSERT OR REPLACE INTO proyectos VALUES (?, ?)', lote)
                    total += len(lote)
                    lote.clear()
            if lote:
                conn.executemany('INSERT OR REPLACE INTO proyectos VALUES (?, ?)', lote)
                total += len(lote)
            conn.execute('INSERT INTO metadata VALUES (?, ?)', ('fuente', origen.name))
            conn.execute('INSERT INTO metadata VALUES (?, ?)', ('filas_leidas', str(total)))
            conn.commit()
        finally:
            conn.close()
        temporal.replace(destino)
        return {'filas_leidas': total, 'db': str(destino), 'columnas': list(indices)}
    finally:
        wb.close()


def consultar_universo(cui: str, db_path: str | Path | None = None) -> dict:
    """Obtiene UNA ficha por CUI sin leer la tabla completa ni clasificaciones."""
    ruta = Path(db_path) if db_path is not None else DB_PATH
    if not ruta.exists():
        return {'encontrado': False, 'cui': str(cui), 'motivo': 'base_121k_no_configurada'}
    with closing(sqlite3.connect(f'file:{ruta.resolve()}?mode=ro', uri=True)) as conn:
        row = conn.execute('SELECT payload FROM proyectos WHERE cui=?', (str(cui),)).fetchone()
    if not row:
        return {'encontrado': False, 'cui': str(cui), 'motivo': 'cui_ausente'}
    datos = json.loads(row[0])
    br = datos.pop('BRECHAS_TODAS', None)
    numero = datos.pop('N_BRECHAS', None)
    return {
        'encontrado': True, 'cui': str(cui), 'fuente': 'universo_2023_2026',
        'proyecto': datos,
        'brechas': [{'descripcion': str(br)[:10000]}] if br not in (None, '') else [],
        'n_brechas': numero,
        'nota': 'Registro del universo filtrado, no clasificacion ni certificacion.',
    }


def consultar_proyecto_v2(cui: str) -> dict:
    """V1 para CUIs historicos; si no existe, consulta universo 2023-2026."""
    from toolbox import obtener_proyecto_dict
    inicial = obtener_proyecto_dict(cui)
    if inicial.get('encontrado'):
        return {**inicial, 'fuente': 'dataset_423_sin_ground_truth'}
    return consultar_universo(cui)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='Preparar/consultar universo 2023-2026')
    parser.add_argument('operacion', choices=['preparar', 'consultar'])
    parser.add_argument('valor', help='Ruta XLSX o CUI')
    parser.add_argument('--db', default=str(DB_PATH))
    args = parser.parse_args()
    salida = preparar_excel(args.valor, args.db) if args.operacion == 'preparar' else consultar_universo(args.valor, args.db)
    print(json.dumps(salida, ensure_ascii=False, indent=2))
